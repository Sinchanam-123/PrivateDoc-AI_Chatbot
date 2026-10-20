import re
import zlib

import pytest

import ingest


# ---------- chunking ----------
def test_chunk_text_empty_and_short():
    assert ingest.chunk_text("") == []
    assert ingest.chunk_text("   \n  ") == []
    assert ingest.chunk_text("A short sentence.") == ["A short sentence."]


def test_chunks_respect_size_and_do_not_cut_words():
    words = [f"word{i}" for i in range(600)]
    text = " ".join(words)
    chunks = ingest.chunk_text(text, size=200, overlap=30)
    assert len(chunks) > 3
    for c in chunks:
        assert len(c) <= 200
    seen = set()
    for c in chunks:
        for w in c.split():
            assert w in words  # no half words
            seen.add(w)
    assert seen == set(words)  # nothing lost


def test_neighbouring_chunks_overlap():
    text = " ".join(f"w{i}" for i in range(300))
    chunks = ingest.chunk_text(text, size=150, overlap=40)
    for a, b in zip(chunks, chunks[1:]):
        assert set(a.split()) & set(b.split())


def test_whitespace_is_normalised():
    assert ingest.chunk_text("a   b\n\n c") == ["a b c"]


# ---------- sections ----------
def test_split_sections_numbers_from_one():
    secs = ingest._split_sections("x" * (ingest.SECTION_CHARS * 2 + 10))
    assert [n for n, _ in secs] == [1, 2, 3]


# ---------- extraction ----------
def test_extract_txt(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello world", encoding="utf-8")
    pages, unit = ingest.extract_pages(f)
    assert unit == "section" and pages == [(1, "hello world")]


def test_extract_docx(tmp_path):
    from docx import Document
    f = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("Second paragraph.")
    d.save(f)
    pages, unit = ingest.extract_pages(f)
    assert unit == "section"
    assert "First paragraph." in pages[0][1] and "Second paragraph." in pages[0][1]


def test_extract_pdf_uses_real_page_numbers(tmp_path):
    f = tmp_path / "a.pdf"
    doc = ingest.fitz.open()
    for text in ("page one text", "page two text"):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(f)
    doc.close()
    pages, unit = ingest.extract_pages(f)
    assert unit == "page"
    assert [n for n, _ in pages] == [1, 2]
    assert "two" in pages[1][1]


def test_unsupported_type_raises(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_text("x")
    with pytest.raises(ValueError):
        ingest.extract_pages(f)


# ---------- indexing end to end (fake embeddings, temporary database) ----------
def fake_embed(texts, prefix):
    out = []
    for t in texts:
        v = [0.0] * 16
        for w in re.findall(r"[a-z0-9]+", (prefix + t).lower()):
            v[zlib.crc32(w.encode()) % 16] += 1.0
        out.append(v)
    return out


@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "CHROMA_DIR", tmp_path / "chroma")
    monkeypatch.setattr(ingest, "DOCS_DIR", tmp_path / "docs")
    monkeypatch.setattr(ingest, "_collection", None)
    monkeypatch.setattr(ingest, "embed", fake_embed)
    (tmp_path / "docs").mkdir()
    return tmp_path / "docs"


def test_index_list_reindex_delete(library):
    f = library / "notes.txt"
    f.write_text("Alpha facts. " * 120, encoding="utf-8")

    n = ingest.index_file(f)
    assert n > 1
    assert ingest.list_sources() == ["notes.txt"]
    count = ingest.get_collection().count()

    ingest.index_file(f)  # re-indexing must not create duplicates
    assert ingest.get_collection().count() == count

    ingest.delete_source("notes.txt")
    assert ingest.list_sources() == []
    assert not f.exists()


def test_file_without_text_indexes_nothing(library):
    f = library / "empty.txt"
    f.write_text("   ", encoding="utf-8")
    assert ingest.index_file(f) == 0


def test_metadata_has_source_and_page(library):
    f = library / "meta.txt"
    f.write_text("Some searchable content here.", encoding="utf-8")
    ingest.index_file(f)
    data = ingest.get_collection().get(include=["metadatas"])
    m = data["metadatas"][0]
    assert m["source"] == "meta.txt" and m["page"] == 1 and m["unit"] == "section"


def test_reindex_all_rebuilds_from_saved_files(library):
    (library / "a.txt").write_text("Alpha content " * 20, encoding="utf-8")
    (library / "b.txt").write_text("Beta content " * 20, encoding="utf-8")
    seen = []
    total = ingest.reindex_all(lambda i, n, name: seen.append((i, n, name)))
    assert total > 0
    assert sorted(ingest.list_sources()) == ["a.txt", "b.txt"]
    assert [s[1] for s in seen] == [2, 2]
